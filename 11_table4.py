"""
11_table4.py
============
Builds manuscript Table 4 (regional causal effects) with both estimands and
95% HDIs, as four sub-tables:

  (a) Fractional anisotropy - total effects        F1-F6
  (b) Fractional anisotropy - controlled direct    F7-F9
  (c) Mean diffusivity - total effects             M1-M5
  (d) Mean diffusivity - controlled direct         M6-M8

Posterior mean, 95% HDI and pd are computed from the beta_exp draws in
Results/{FA,MD}_B/*_idata.nc. Cells read "+0.34 [0.18; 0.49]" and are bold
when pd >= 0.975.

Outputs Table4_regional_effects.docx (Calibri 10 pt, portrait) and a .md twin.

Usage:
  python 11_table4.py
"""

from __future__ import annotations

import argparse
import logging
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from paper_style import EXPOSURE_SHORT, ROI_ORDER, load_regional_table, rlabel

logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
log = logging.getLogger(__name__)

BASE_DIR = Path(__file__).parent
RESULTS_DIR = BASE_DIR / "Results"

FONT_NAME = "Calibri"
FONT_SIZE_PT = 10
PD_BOLD_THRESHOLD = 0.975

# A4 portrait with 2 cm margins -> 17.0 cm of usable width.
PAGE_WIDTH_CM = 21.0
PAGE_HEIGHT_CM = 29.7
MARGIN_CM = 2.0
USABLE_WIDTH_CM = PAGE_WIDTH_CM - 2 * MARGIN_CM
ROI_COL_CM = 2.0

SUBTABLES = [
    ("a", "FA", "Fractional anisotropy - total effects", ["F1", "F2", "F3", "F4", "F5", "F6"]),
    ("b", "FA", "Fractional anisotropy - controlled direct effects", ["F7", "F8", "F9"]),
    ("c", "MD", "Mean diffusivity - total effects", ["M1", "M2", "M3", "M4", "M5"]),
    ("d", "MD", "Mean diffusivity - controlled direct effects", ["M6", "M7", "M8"]),
]

CAPTION = (
    "Table 4. Regional causal effects (posterior mean beta on the z-scored scale, "
    "95% highest-density interval in brackets). Sub-tables (a) and (c) give total "
    "effects; (b) and (d) give controlled direct effects, which are identified only "
    "for axon density, astrocyte density and microglia density. Values in bold have "
    "pd >= 0.975. Full per-estimand tables including pd are given in the Supplement "
    "(Tables S2-S3)."
)


def format_cell(row) -> tuple[str, bool]:
    text = f"{row['mean']:+.2f} [{row['hdi_2.5']:.2f}; {row['hdi_97.5']:.2f}]"
    p_gt0 = float(row["P(>0)"])
    return text, max(p_gt0, 1.0 - p_gt0) >= PD_BOLD_THRESHOLD


def build_matrix(table, models: list[str]) -> dict[tuple[str, str], tuple[str, bool]]:
    cells: dict[tuple[str, str], tuple[str, bool]] = {}
    for model_id in models:
        sub = table[table["model"] == model_id].set_index("roi_type")
        for roi in ROI_ORDER:
            cells[(roi, model_id)] = format_cell(sub.loc[roi]) if roi in sub.index else ("-", False)
    return cells


def style_run(run, bold: bool) -> None:
    run.font.name = FONT_NAME
    run.font.size = Pt(FONT_SIZE_PT)
    run.bold = bold


def add_docx_subtable(doc: Document, letter: str, title: str, models: list[str], cells) -> None:
    heading = doc.add_paragraph()
    style_run(heading.add_run(f"({letter}) {title}"), bold=True)

    table = doc.add_table(rows=1 + len(ROI_ORDER), cols=1 + len(models))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    value_col_cm = (USABLE_WIDTH_CM - ROI_COL_CM) / len(models)

    def write(cell, text: str, bold: bool, centered: bool) -> None:
        paragraph = cell.paragraphs[0]
        if centered:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_run(paragraph.add_run(text), bold)

    header = table.rows[0].cells
    write(header[0], "ROI type", True, False)
    for col, model_id in enumerate(models, start=1):
        write(header[col], EXPOSURE_SHORT[model_id].replace(" CDE", ""), True, True)

    for row_idx, roi in enumerate(ROI_ORDER, start=1):
        cell_row = table.rows[row_idx].cells
        write(cell_row[0], rlabel(roi), False, False)
        for col, model_id in enumerate(models, start=1):
            text, bold = cells[(roi, model_id)]
            write(cell_row[col], text, bold, True)

    # Column widths must be set on every cell for Word to honour them.
    for row in table.rows:
        row.cells[0].width = Cm(ROI_COL_CM)
        for cell in row.cells[1:]:
            cell.width = Cm(value_col_cm)

    doc.add_paragraph()


def markdown_subtable(letter: str, title: str, models: list[str], cells) -> str:
    headers = ["ROI type"] + [EXPOSURE_SHORT[m].replace(" CDE", "") for m in models]
    lines = [f"**({letter}) {title}**", "", "| " + " | ".join(headers) + " |",
             "|" + "|".join(["---"] * len(headers)) + "|"]
    for roi in ROI_ORDER:
        row = [rlabel(roi)]
        for model_id in models:
            text, bold = cells[(roi, model_id)]
            row.append(f"**{text}**" if bold else text)
        lines.append("| " + " | ".join(row) + " |")
    lines.append("")
    return "\n".join(lines)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build manuscript Table 4.")
    parser.add_argument("--results-dir", type=Path, default=RESULTS_DIR)
    parser.add_argument("--out-stem", type=Path, default=BASE_DIR / "Table4_regional_effects")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    tables = {outcome: load_regional_table(args.results_dir, outcome) for outcome in ("FA", "MD")}

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(FONT_SIZE_PT)

    section = doc.sections[0]
    section.page_width = Cm(PAGE_WIDTH_CM)
    section.page_height = Cm(PAGE_HEIGHT_CM)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(MARGIN_CM))

    caption = doc.add_paragraph()
    style_run(caption.add_run(CAPTION), bold=False)
    doc.add_paragraph()

    md_parts = [CAPTION, ""]
    for letter, outcome, title, models in SUBTABLES:
        cells = build_matrix(tables[outcome], models)
        add_docx_subtable(doc, letter, title, models, cells)
        md_parts.append(markdown_subtable(letter, title, models, cells))

    docx_path = args.out_stem.with_suffix(".docx")
    md_path = args.out_stem.with_suffix(".md")
    doc.save(docx_path)
    md_path.write_text("\n".join(md_parts))
    log.info("Saved %s", docx_path)
    log.info("Saved %s", md_path)


if __name__ == "__main__":
    main()
